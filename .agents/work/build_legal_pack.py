from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path('outputs/Platform_Legal_and_Policy_Draft_Pack.docx')
BLUE = RGBColor(75, 46, 232)
DARK = RGBColor(25, 21, 40)
MUTED = RGBColor(104, 99, 116)
LIGHT = 'F2F4F7'
LAV = 'EEEAFE'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd')) or OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    if shd.getparent() is None: tcPr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar') or OxmlElement('w:tcMar')
    if tcMar.getparent() is None: tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}')) or OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')
        if node.getparent() is None: tcMar.append(node)

def mark_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    trPr.append(header)

def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(total)); tblW.set(qn('w:type'),'dxa')
    tblInd = tblPr.find(qn('w:tblInd'))
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'),'120'); tblInd.set(qn('w:type'),'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'),str(width)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'),str(widths[idx])); tcW.set(qn('w:type'),'dxa')
            set_cell_margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_font(run, size=11, bold=False, color=DARK, italic=False):
    run.font.name='Calibri'; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=color

def add_page_field(paragraph):
    r = paragraph.add_run(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)

doc = Document()
sec = doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.font.color.rgb=DARK
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
for name,size,color,before,after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',12,RGBColor(31,77,120),8,4)]:
    st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=color
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
for sname in ['List Bullet','List Number','List Number 2']:
    st=styles[sname]; st.font.name='Calibri'; st.font.size=Pt(11); st.paragraph_format.left_indent=Inches(.5); st.paragraph_format.first_line_indent=Inches(-.25); st.paragraph_format.space_after=Pt(8); st.paragraph_format.line_spacing=1.167
if 'Callout' not in styles:
    c=styles.add_style('Callout',WD_STYLE_TYPE.PARAGRAPH); c.font.name='Calibri'; c.font.size=Pt(10.5); c.font.color.rgb=DARK; c.paragraph_format.space_before=Pt(6); c.paragraph_format.space_after=Pt(10); c.paragraph_format.left_indent=Inches(.18); c.paragraph_format.right_indent=Inches(.18)

header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run('PLATFORM GOVERNANCE DRAFT PACK  |  CONFIDENTIAL'),8,True,MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run('Lawyer review draft  |  Page '),8,False,MUTED); add_page_field(footer)

def title(text, subtitle=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(32); p.paragraph_format.space_after=Pt(6)
    set_font(p.add_run(text),28,True,DARK)
    if subtitle:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(24); set_font(p.add_run(subtitle),13,False,MUTED)

def para(text='', bold_lead=None, style=None):
    p=doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead),11,True); set_font(p.add_run(text[len(bold_lead):]),11)
    else: set_font(p.add_run(text),11)
    return p

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); set_font(p.add_run(text),11); return p

def number(text):
    p=doc.add_paragraph(style='List Number'); set_font(p.add_run(text),11); return p

def number_restart(text):
    p=doc.add_paragraph(style='List Number 2'); set_font(p.add_run(text),11); return p

def callout(label,text):
    p=doc.add_paragraph(style='Callout')
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'),LAV); pPr.append(shd)
    set_font(p.add_run(label+'  '),10.5,True,BLUE); set_font(p.add_run(text),10.5)

def h1(text): doc.add_heading(text,level=1)
def h2(text): doc.add_heading(text,level=2)
def h3(text): doc.add_heading(text,level=3)
def pagebreak(): doc.add_page_break()

title('Platform Legal & Policy Draft Pack','Working brand: Potchly / successor national brand  |  Prepared 30 June 2026')
callout('IMPORTANT','This pack is a structured first draft for review by a South African attorney. It is not legal advice and must not be published until company details, vendor arrangements, commercial rules and statutory requirements have been confirmed.')
h2('Document-control placeholders')
for item in ['[LEGAL COMPANY NAME]','[COMPANY REGISTRATION NUMBER]','[REGISTERED AND PHYSICAL ADDRESS]','[PLATFORM NAME AND DOMAIN]','[INFORMATION OFFICER NAME AND CONTACT]','[SUPPORT AND LEGAL EMAILS]','[EFFECTIVE DATE]','[COMMISSION, FEES AND CANCELLATION WINDOWS]']:
    bullet(item)
h2('Contents')
for item in ['A. Privacy Notice','B. Customer Terms of Use','C. Provider Agreement','D. Cancellation and Refund Policy','E. Dispute Resolution Policy','F. Acceptable Use and Reviews Policy','G. Data Retention Schedule','H. Security Incident Response Plan','I. Vendor Data-Processing Checklist','J. PAIA Manual Working Draft','K. Lawyer Review Questions and Sources']:
    bullet(item)

pagebreak(); h1('A. Privacy Notice')
para('Effective date: [EFFECTIVE DATE]. Responsible party: [LEGAL COMPANY NAME] trading as [PLATFORM NAME] ("we", "us" or "Platform").')
h2('1. Scope and role')
para('This notice explains how we process personal information when customers, service providers, visitors, employees or authorised representatives use our website, applications, communications and marketplace services. We act as the responsible party for Platform account and marketplace information. Payment and identity-verification partners may act as independent responsible parties or operators, depending on their contracts and processing activities.')
h2('2. Information we collect')
for text in ['Account information: name, contact details, authentication records, language, suburb and preferences.','Provider information: identity-verification result, business details, service categories, portfolio, qualifications, service area, payout-account verification result and compliance status.','Marketplace records: requests, bids, bookings, messages, reviews, disputes, evidence, cancellations and support communications.','Payment information: transaction references, amount, status, refunds and payout status. Card and bank-login credentials are processed on the payment partner’s hosted service and should not be stored by the Platform.','Device and security information: IP address, browser/device data, login history, risk signals, cookie choices and audit events.','Location information: approximate service area before booking and, where necessary, an exact address disclosed only for a confirmed booking.','Special personal information: biometric or identity information is processed only with an appropriate lawful basis and safeguards, preferably directly by the verification partner.']:
    bullet(text)
h2('3. Why and on what basis we process information')
for text in ['Create and secure accounts; verify contact details; prevent fraud and account takeover.','Match customers and providers; enable bids, bookings, messaging and reviews.','Administer protected payments, payouts, refunds, chargebacks and disputes.','Verify providers and protect users, where consent, contract, legal obligation or another lawful justification applies.','Comply with legal, tax, accounting, regulatory and law-enforcement obligations.','Send service messages. Marketing is optional and may be withdrawn without affecting service notifications.','Improve reliability and safety using aggregated or de-identified information where reasonably possible.']:
    bullet(text)
h2('4. Sharing and operators')
para('We disclose only information reasonably required for the purpose. Recipients may include the selected customer or provider; payment, identity, messaging, hosting, security and professional-service providers; regulators; courts; law-enforcement bodies; insurers; and a successor in a lawful corporate transaction. Contracts with operators must address confidentiality, security, breach notification, retention, deletion and sub-processing.')
h2('5. Identity checks and automated decisions')
para('Provider verification may include document authenticity, government or trusted-source checks, facial comparison and liveness detection. The Platform should retain the outcome and reference rather than a reusable copy of biometric material where feasible. A provider may request human review if an automated result prevents onboarding or materially affects access.')
h2('6. Payments')
para('Payment credentials are entered on the payment partner’s secure interface. We retain marketplace transaction references and statuses needed to administer bookings, refunds, disputes and accounting. The applicable payment partner’s own privacy notice also applies.')
h2('7. Cross-border processing')
para('If information is processed outside South Africa, we will document the transfer, recipient, location, safeguards and lawful basis, and require protections consistent with applicable law. The lawyer must verify the final hosting and vendor locations before publication.')
h2('8. Retention and deletion')
para('We retain information only for as long as necessary for the stated purpose, contract, safety need or legal obligation, then delete, de-identify or securely archive it. Indicative periods appear in the Data Retention Schedule and must be validated by legal and tax advisers.')
h2('9. Security')
para('Safeguards include encryption in transit and at rest, least-privilege access, multi-factor authentication for privileged users, secure development, audit logging, backups, monitoring, vendor due diligence and incident procedures. No system is completely secure; users must protect their credentials and report suspicious activity promptly.')
h2('10. Your choices and rights')
for text in ['Ask whether we hold personal information and request access, correction or deletion where applicable.','Object to processing or withdraw consent where consent is the basis.','Opt out of direct marketing.','Request information about recipients or automated decisions where applicable.','Complain to our Information Officer or the Information Regulator.']:
    bullet(text)
h2('11. Children')
para('The Platform is intended for persons aged 18 or older. We do not knowingly onboard children as customers or providers without a verified lawful basis and competent-person involvement approved by counsel.')
h2('12. Contact')
para('Information Officer: [NAME], [EMAIL], [PHONE], [ADDRESS]. Information Regulator: see https://inforegulator.org.za/.')

pagebreak(); h1('B. Customer Terms of Use')
h2('1. Agreement and eligibility')
para('By creating an account or using the Platform, the customer accepts these terms, the Privacy Notice and incorporated policies. Customers must be at least 18, legally capable, and provide accurate information. Acceptance must be recorded with the version and timestamp.')
h2('2. Platform role')
para('The Platform facilitates introductions, bids, bookings, protected payments and dispute administration. Unless expressly stated, providers are independent businesses and not employees or agents of the Platform. The final service contract is between the customer and selected provider, subject to consumer law and the Platform protection rules.')
h2('3. Requests, bids and bookings')
for text in ['A request must be lawful, accurate and sufficiently detailed.','A bid states price, scope, date, location basis, inclusions, exclusions and expiry.','A booking is created when the customer accepts a bid and completes the required payment or deposit.','Material changes require recorded agreement; providers may not substitute undisclosed persons without permission.']:
    bullet(text)
h2('4. Pricing and payment protection')
para('The checkout must disclose the service price, Platform fee, taxes if applicable, cancellation terms and total before payment. Funds are processed and, where applicable, held or split by the payment partner. The Platform must not describe funds as escrow unless the contracted provider legally supplies that service.')
h2('5. Completion, release and disputes')
para('The customer should confirm completion only after the service has been supplied. Release may occur automatically after a disclosed period if no problem is reported. A timely dispute pauses payout where the payment arrangement permits. Fraudulent disputes or collusion may lead to suspension and recovery action.')
h2('6. Cancellations and refunds')
para('The Cancellation and Refund Policy forms part of these terms. Statutory rights prevail. The final policy must distinguish provider cancellation, customer cancellation, non-attendance, unsafe conduct, defective service and force majeure.')
h2('7. Customer conduct and safety')
for text in ['Treat providers lawfully and respectfully; provide safe access and accurate instructions.','Do not solicit off-platform payment for a booking introduced by the Platform during the protected period stated in the final terms.','Do not upload unlawful, discriminatory, threatening, sexually explicit, fraudulent or privacy-invasive content.','For immediate danger contact emergency services; the Platform is not an emergency service.']:
    bullet(text)
h2('8. Reviews')
para('Reviews must reflect a genuine booking and honest experience. We may remove content that violates the Reviews Policy, but will not manipulate legitimate criticism. Providers should have a fair response process.')
h2('9. Suspension and termination')
para('We may restrict or suspend an account for safety, fraud, unlawful conduct, repeated cancellations, payment abuse or policy breaches. Except where immediate action is necessary, notice and a reasonable opportunity to respond should be provided. Users may close their account subject to unresolved bookings, disputes and lawful retention.')
h2('10. Liability and consumer rights')
para('No clause excludes liability or rights that cannot lawfully be excluded. Counsel must draft balanced limitations for indirect loss, Platform availability and independent provider conduct, taking account of the Consumer Protection Act and the actual operational model.')
h2('11. Governing law and complaints')
para('South African law applies. Complaints should first follow the Platform dispute process without limiting any right to approach an ombud, regulator, the National Consumer Commission, tribunal or court with jurisdiction.')

pagebreak(); h1('C. Provider Agreement')
h2('1. Provider status and onboarding')
para('The provider is an independent business responsible for its services, personnel, permits, qualifications, tax and insurance. No employment, partnership, agency or exclusivity is created. Onboarding remains conditional on contact, identity, bank, portfolio and risk checks.')
h2('2. Accurate profile and authorised personnel')
para('The provider must maintain accurate identity, business, service, pricing, location and availability information. Only approved personnel may fulfil bookings. Verification badges indicate completed checks, not a guarantee of skill, character or future performance.')
h2('3. Bids and service standards')
for text in ['Bid only where capable, available and properly equipped.','State the total price and all material inclusions, exclusions, travel fees, products and time assumptions.','Supply services with reasonable care, skill, safety, punctuality and lawful products.','Protect customer information and use it only for the booking.','Record agreed scope changes through the Platform.']:
    bullet(text)
h2('4. Fees, settlement and tax')
para('The provider authorises the disclosed Platform commission and payment-partner deductions. Settlement timing is subject to completion, dispute, refund, reserve, chargeback and compliance rules. The provider remains responsible for tax invoices, income declarations and taxes unless law requires the Platform to withhold or report.')
h2('5. Cancellations, refunds and re-performance')
para('Provider cancellation, non-attendance, unsafe conduct, material misrepresentation or materially defective service may result in refund, re-performance, reduced payment, fee recovery, ranking impact or suspension, subject to law and fair review.')
h2('6. Data, confidentiality and communication')
para('Customer contact, address and booking information is confidential and may not be copied, marketed to, sold or retained beyond the permitted purpose. Security incidents must be reported immediately to [SECURITY EMAIL]. Off-platform contact does not remove these obligations.')
h2('7. Intellectual property and content')
para('The provider retains ownership of original portfolio content but grants the Platform a non-exclusive, revocable licence to display it for operating and promoting the marketplace. The provider warrants that uploads are authorised and do not infringe rights.')
h2('8. Monitoring, suspension and appeals')
para('The Platform may investigate risk signals, complaints and performance. Urgent restrictions may be imposed to protect users or funds. The provider should receive reasons and a human appeal route unless disclosure would compromise security, legal duties or another person’s rights.')
h2('9. Insurance and indemnity')
para('Counsel must decide which categories require public-liability, professional-indemnity or product insurance and appropriate minimum limits. Any indemnity must be lawful, proportionate and consistent with consumer protection and the Platform’s own responsibilities.')
h2('10. Termination')
para('Either party may terminate on [NOTICE PERIOD], subject to existing bookings, disputes, payouts and lawful retention. The Platform may terminate immediately for fraud, serious safety risk, unlawful conduct or repeated material breach.')

pagebreak(); h1('D. Cancellation and Refund Policy')
para('The table below is a proposed operating position, not final law. Counsel and the selected payment provider must approve the windows, fees and release mechanics.')
t=doc.add_table(rows=1,cols=4); set_table_geometry(t,[1800,2200,2600,2760]); mark_header(t.rows[0])
for i,h in enumerate(['Scenario','Default outcome','Payment treatment','Evidence / notes']): set_cell_shading(t.cell(0,i),LIGHT); set_font(t.cell(0,i).paragraphs[0].add_run(h),9.5,True)
rows=[
('Provider cancels','Full customer refund','Provider receives no payout; Platform fee refunded unless unavoidable external cost is lawfully disclosed','Reason, timestamp and communications recorded'),
('Customer cancels early','Full or near-full refund','Only unavoidable, disclosed costs may be retained','Final early-cancellation window: [X hours]'),
('Customer cancels late','Partial refund may apply','Reasonable provider compensation only if disclosed and lawful','Final late window and cap: [INSERT]'),
('Provider no-show','Full refund plus risk review','Payout remains blocked','Attendance and message records'),
('Customer no-show','Provider may receive disclosed portion','Platform applies fair grace period and contact attempts','Address privacy and safety considered'),
('Materially defective service','Re-performance, partial or full refund','Payout paused pending resolution','Photos, scope, offer and communications'),
('Safety incident','Immediate pause and investigation','No release until safety review','Emergency channels take priority'),
('Force majeure','Fair allocation case by case','Avoid penalties where neither party is at fault','Document event and mitigation')]
for row in rows:
    cells=t.add_row().cells
    for i,val in enumerate(row): set_font(cells[i].paragraphs[0].add_run(val),9.2)
set_table_geometry(t,[1800,2200,2600,2760])
h2('Refund timing')
para('Approved refunds are submitted promptly to the payment partner. Bank and card processing times are outside the Platform’s direct control. The user receives a reference and status updates.')

pagebreak(); h1('E. Dispute Resolution Policy')
h2('1. Objectives')
para('Resolve complaints fairly, quickly and consistently; preserve evidence; protect funds where contractually possible; give both parties a chance to respond; identify safety patterns; and preserve external legal rights.')
h2('2. Process')
for text in ['User opens a case from the booking and selects the issue.','System captures the accepted bid, booking, payment events, messages and submitted evidence.','Where permitted, payout is paused and both parties are notified.','The responding party receives [2 business days] unless urgency requires less.','A trained reviewer assesses scope, conduct, evidence, consumer rights and previous relevant incidents.','Outcome may include release, re-performance, partial refund, full refund, credit, warning, suspension or external referral.','Both parties receive reasons and an internal appeal route within [5 business days].']:
    number(text)
h2('3. Safety and criminal allegations')
para('Immediate danger must be directed to emergency services. The Platform may preserve records, restrict accounts and cooperate with lawful requests. Staff must not promise confidentiality that cannot be maintained or conduct amateur criminal investigations.')
h2('4. Independence and conflicts')
para('Reviewers must disclose conflicts and may not decide cases involving themselves, close associates or material interests. High-value, safety-related or precedent-setting cases require senior review.')

pagebreak(); h1('F. Acceptable Use and Reviews Policy')
h2('Prohibited conduct')
for text in ['Fraud, impersonation, account sharing, payment circumvention or manipulation of bids.','Threats, harassment, discrimination, hate content, exploitation or unsafe sexual conduct.','Unlawful services, controlled goods, stolen property or services requiring unauthorised professional practice.','Malware, credential harvesting, scraping, interference, false reports or abuse of support.','Publishing another person’s address, identity number, financial data or private communications without lawful justification.','Fake reviews, review trading, coercion or retaliation for honest feedback.']:
    bullet(text)
h2('Content moderation and appeals')
para('The Platform may remove, limit or preserve content based on law, safety and policy. Decisions should be proportionate and recorded. Users receive a reason and appeal route where appropriate. Evidence relevant to disputes or legal duties may be retained even if removed from public view.')

pagebreak(); h1('G. Data Retention Schedule')
para('These are proposed business periods. Counsel, the accountant, insurers and vendors must confirm statutory and contractual requirements before adoption.')
t=doc.add_table(rows=1,cols=4); set_table_geometry(t,[1900,2600,1700,3160]); mark_header(t.rows[0])
for i,h in enumerate(['Record','Purpose','Proposed period','End-of-period action']): set_cell_shading(t.cell(0,i),LIGHT); set_font(t.cell(0,i).paragraphs[0].add_run(h),9.5,True)
rows=[
('Unverified signup','Complete registration / prevent abuse','30 days','Delete or de-identify'),('Active account','Provide services and security','Account life','Review on closure'),('Closed account profile','Claims, fraud and reactivation control','3 years, subject to advice','Delete or de-identify'),('Booking and transaction ledger','Contract, accounting, tax and disputes','5 years or statutory period','Restricted archive then delete'),('Messages linked to booking','Safety and dispute evidence','3 years after closure','Delete or de-identify'),('Exact service address','Fulfil confirmed booking','90 days after completion unless disputed','Delete precise value; retain suburb if needed'),('Raw identity/biometric material','Vendor verification','Prefer no Platform storage','Vendor deletes under contract'),('Verification result/reference','Provider trust and audit','Provider life + 3 years','Restricted archive then delete'),('Dispute evidence','Resolve and defend claim','5 years after closure','Delete securely'),('Security logs','Detect and investigate abuse','12 months; longer for incidents','Rotate/delete'),('Marketing consent','Prove choice and preferences','Until withdrawal + evidence period','Suppress and retain minimal proof'),('Backups','Recovery','35-90 days rolling','Automatic encrypted expiry')]
for row in rows:
    cells=t.add_row().cells
    for i,val in enumerate(row): set_font(cells[i].paragraphs[0].add_run(val),9.1)
set_table_geometry(t,[1900,2600,1700,3160])

pagebreak(); h1('H. Security Incident Response Plan')
h2('1. Roles')
para('Incident lead: [NAME]. Information Officer: [NAME]. Technical lead: [NAME]. Legal adviser: [FIRM]. Vendor contacts: [LIST]. Only authorised spokespeople communicate externally.')
h2('2. Response sequence')
for text in ['Detect and log: create a timestamped incident record; preserve evidence.','Triage: determine affected systems, data, users, vendors and ongoing risk.','Contain: revoke credentials, isolate services, block abuse and protect funds without destroying evidence.','Assess: confirm categories and volume of information, likely harm and cross-border/vendor impact.','Notify: obtain legal advice and notify the Information Regulator and affected people as required, without unreasonable delay.','Recover: restore clean systems, rotate secrets, monitor recurrence and support affected users.','Review: document root cause, decisions, timeline, lessons and accountable remediation.']:
    number_restart(text)
h2('3. Minimum preparation')
for text in ['24/7 escalation list and vendor breach contacts.','Current asset, data-flow and subprocessor inventories.','Tested backups and restoration exercises.','Prepared notification templates and identity-theft support options.','Quarterly tabletop exercise and annual independent security review.']:
    bullet(text)

pagebreak(); h1('I. Vendor Data-Processing and Contract Checklist')
for heading,items in [
('Commercial',['Setup, minimum, transaction, verification, messaging, refund and payout fees','Settlement timing, reserves, chargebacks and termination costs','Sandbox, support, uptime and incident service levels']),
('Data protection',['Role: operator or independent responsible party','Purpose limitation, documented instructions and confidentiality','Hosting locations, cross-border safeguards and subprocessor approval','Retention, deletion, return, audit evidence and data-subject assistance','Security controls and breach notice measured in hours, not merely “promptly”']),
('Technical',['API/SDK maturity, webhook signing and idempotency','Separate test and production keys; rotation and least privilege','Error handling, reconciliation, export and business continuity']),
('Risk and exit',['Liability, insurance, warranties and regulatory standing','No broad right to use marketplace data for unrelated marketing','Orderly transition, data export and deletion certificate'])]:
    h2(heading)
    for item in items: bullet(item)

pagebreak(); h1('J. PAIA Manual Working Draft')
callout('COMPLETE FROM THE OFFICIAL TEMPLATE','This section is a working skeleton. The final manual should follow the Information Regulator’s current Private Body template and include the prescribed forms and availability arrangements.')
h2('1. Body and Information Officer')
para('Private body: [LEGAL COMPANY NAME], registration [NUMBER], addresses [DETAILS]. Head / Information Officer: [DETAILS]. Deputy Information Officer(s): [DETAILS].')
h2('2. Guide and request process')
para('The Information Regulator’s PAIA Guide is available from the Regulator. Requests to this body must use the prescribed form, identify the record and right involved, provide contact and identity information, state the preferred access method and pay any prescribed fee. The Information Officer will communicate the outcome and available remedies.')
h2('3. Records available without a formal request')
for text in ['Published privacy notices, policies and terms','Public provider profiles and service categories','Company and regulatory information made available on the website','Marketing and media material intended for public distribution']:
    bullet(text)
h2('4. Record categories')
for text in ['Company, governance and statutory records','Financial, tax, banking, payment and accounting records','Customer, provider, booking, communication and dispute records','Personnel, contractor and recruitment records','Vendor, procurement, insurance and professional-adviser records','Security, access-control, audit, incident and business-continuity records','Intellectual property, software, domain and brand records']:
    bullet(text)
h2('5. Grounds, remedies and availability')
para('Access may be granted, refused, redacted or deferred only in accordance with applicable law. The final manual must specify internal contacts, complaint and court remedies, fees, languages, office availability, website publication and the accessibility of facilities to people with disabilities.')

pagebreak(); h1('K. Lawyer Review Questions and Sources')
h2('Decisions requiring counsel')
for text in ['Is the Platform legally an intermediary, supplier, agent or combination for each service and payment flow?','Which CPA cancellation, quality, disclosure and refund duties apply, and which terms would be prohibited?','May the selected payment model be described as escrow, and who is responsible for refunds and chargebacks?','What provider insurance, licences and category exclusions are required?','What lawful basis and prior-authorisation questions apply to biometrics, criminal allegations or other special personal information?','What cross-border safeguards are required for hosting, support and analytics?','Which record periods are mandatory for tax, company, payment and litigation purposes?','Does any automated ranking or verification decision require additional notice, human review or prior authorisation?','Which PAIA records and forms must be published, and in which languages?']:
    bullet(text)
h2('Primary reference sources')
for text in ['Protection of Personal Information Act 4 of 2013: https://www.gov.za/documents/protection-personal-information-act','Information Regulator POPIA and PAIA guidance: https://inforegulator.org.za/','Consumer Protection Act 68 of 2008: https://www.gov.za/documents/consumer-protection-act','Electronic Communications and Transactions Act 25 of 2002: https://www.gov.za/documents/electronic-communications-and-transactions-act','National Consumer Commission: https://thencc.org.za/']:
    bullet(text)

doc.core_properties.title='Platform Legal and Policy Draft Pack'
doc.core_properties.subject='Lawyer review draft for South African local-services marketplace'
doc.core_properties.author='[LEGAL COMPANY NAME]'
doc.save(OUT)
print(OUT.resolve())
